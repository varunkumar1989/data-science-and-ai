"""SkillNavora Data Science & AI Handbook builder.

python3 build.py            -> runs every lesson (cached), writes course/lessons/*.html, course/index.html, course/print/*.html, pdf_jobs.json
node make_pdfs.js           -> prints the PDFs listed in pdf_jobs.json

Lesson sources live in source/lessons/<track>/<NN>-<slug>.md (Markdown + front matter). Special blocks:
  ```python run``` / ```python fig: Caption```   executed with real data, output and charts captured
  ```sql run```                                    executed (T-SQL is translated for the local engine), result table shown
  ```xl [title="..."] [formulas]```               spreadsheet grid; formulas really calculated (LibreOffice)
  :::tip|note|warn|example|live|new|interview|key|mistake|ai Title ... :::     callouts
  :::quiz ... :::  (Q: / - option / * correct / > explanation)      :::exercise Title ... ??? solution ... :::
"""
import os, re, sys, json, glob, html, hashlib, ast, io, contextlib, traceback, shutil, subprocess, datetime, warnings, textwrap
import yaml, markdown
import pandas as pd, numpy as np

warnings.filterwarnings("ignore")
SRC = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(SRC)
DATA = os.path.join(ROOT, "datasets"); IMG = os.path.join(ROOT, "img"); LES = os.path.join(ROOT, "lessons"); PRINT = os.path.join(ROOT, "print")
FILES = os.path.join(ROOT, "files"); CACHE = os.path.join(SRC, ".cache"); PDFDIR = os.path.join(ROOT, "pdf")
for d in (IMG, LES, PRINT, FILES, CACHE, PDFDIR): os.makedirs(d, exist_ok=True)
sys.path.insert(0, SRC)
import dg  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402

TITLE = "Data Science & AI Handbook"
BRAND = "SkillNavora"
AUTHOR = "Varun Kumar"
REPO = "https://github.com/varunkumar1989/data-science-and-ai"
SITE = "https://www.skillnavora.in"
UPDATED = "October 2026"
VERSION = "2026.10"
ARTIFACT = os.environ.get("HB_ARTIFACT", "").strip()

W3 = "https://www.w3schools.com"
TRACKS = [
    dict(id="excel", name="Advanced Excel", short="Excel", color="#1E8E5A", desc="From the Excel window to dashboards: formatting, 100+ functions, lookups, pivot tables, advanced charts, Power Query, Power Pivot and AI in Excel.",
         practice=[("W3Schools Excel tutorial", f"{W3}/excel/index.php"), ("W3Schools exercises hub (Excel)", f"{W3}/exercises/"), ("Microsoft Excel training (official)", "https://support.microsoft.com/en-us/excel"),
                   ("Excel functions reference (Microsoft)", "https://support.microsoft.com/en-us/office/excel-functions-alphabetical-b3944572-255d-4efb-bb96-c6d90033e188"), ("Chandoo.org free Excel lessons", "https://chandoo.org/wp/excel-school/")]),
    dict(id="sql", name="SQL for Data Analysis", short="SQL", color="#2563EB", desc="Relational databases, ER diagrams and keys, DDL/DML, filtering, aggregation, joins, set operators, subqueries, window functions, CTEs, views, procedures and performance.",
         practice=[("W3Schools SQL tutorial", f"{W3}/sql/default.asp"), ("W3Schools SQL exercises", f"{W3}/sql/sql_exercises.asp"), ("W3Schools SQL quiz", f"{W3}/quiztest/quiztest.asp?qtest=SQL"),
                   ("SQLBolt interactive lessons", "https://sqlbolt.com/"), ("LeetCode SQL 50 study plan", "https://leetcode.com/studyplan/top-sql-50/"), ("HackerRank SQL practice", "https://www.hackerrank.com/domains/sql"),
                   ("T-SQL reference (Microsoft Learn)", "https://learn.microsoft.com/en-us/sql/t-sql/language-reference")]),
    dict(id="powerbi", name="Power BI", short="Power BI", color="#E2A400", desc="ETL with Power Query, star-schema modelling, DAX measures and time intelligence, interactive reports, case studies, publishing, security and Copilot.",
         practice=[("Microsoft Learn: Power BI data analyst path (PL-300)", "https://learn.microsoft.com/en-us/training/career-paths/data-analyst"), ("DAX function reference", "https://learn.microsoft.com/en-us/dax/dax-function-reference"),
                   ("DAX Guide", "https://dax.guide/"), ("SQLBI articles and videos", "https://www.sqlbi.com/"), ("Power BI sample datasets (Microsoft)", "https://learn.microsoft.com/en-us/power-bi/create-reports/sample-datasets")]),
    dict(id="ai-analytics", name="AI for Data Analysis", short="AI Analytics", color="#DB2777", desc="Generative AI for analysts: prompting, AI vs GenAI vs agentic AI, AI over SQL (Metabase, MinusX), Copilot and Gemini in Excel, and office agents.",
         practice=[("W3Schools Generative AI tutorial", f"{W3}/gen_ai/"), ("Microsoft Copilot in Excel help", "https://support.microsoft.com/en-us/copilot-excel"), ("Metabase docs", "https://www.metabase.com/docs/latest/"),
                   ("Anthropic prompt engineering guide", "https://docs.claude.com/en/docs/build-with-claude/prompt-engineering/overview"), ("Google Prompting guide for Gemini", "https://ai.google.dev/gemini-api/docs/prompting-strategies")]),
    dict(id="python", name="Python for Data Science", short="Python", color="#7C3AED", desc="Python from zero: data types, control flow, data structures, functions, OOP, files, NumPy, pandas, visualisation, dates, SQL connections, EDA and DSA for data roles.",
         practice=[("W3Schools Python tutorial", f"{W3}/python/default.asp"), ("W3Schools Python exercises", f"{W3}/python/python_exercises.asp"), ("W3Schools Python quiz", f"{W3}/quiztest/quiztest.asp?qtest=PYTHON"),
                   ("W3Schools NumPy exercises", f"{W3}/python/numpy/numpy_exercises.asp"), ("W3Schools Pandas exercises", f"{W3}/python/pandas/pandas_exercises.asp"), ("W3Schools Pandas quiz", f"{W3}/quiztest/quiztest.asp?qtest=PANDAS"),
                   ("W3Schools DSA tutorial", f"{W3}/dsa/index.php"), ("Kaggle Learn: Python and Pandas (free)", "https://www.kaggle.com/learn"), ("pandas user guide", "https://pandas.pydata.org/docs/user_guide/index.html")]),
    dict(id="ml", name="Statistics & Machine Learning", short="Stats & ML", color="#0E9F8E", desc="Probability, CLT, hypothesis tests, the ML lifecycle, regression, classification, trees, ensembles, clustering, PCA, time series, NLP, deep learning basics, evaluation, explainability and deployment.",
         practice=[("W3Schools Statistics tutorial", f"{W3}/statistics/index.php"), ("W3Schools Python ML tutorial", f"{W3}/python/python_ml_getting_started.asp"), ("W3Schools Data Science tutorial", f"{W3}/datascience/default.asp"),
                   ("scikit-learn user guide", "https://scikit-learn.org/stable/user_guide.html"), ("Kaggle Learn: Intro & Intermediate ML", "https://www.kaggle.com/learn"), ("Google ML Crash Course", "https://developers.google.com/machine-learning/crash-course"),
                   ("StatQuest (YouTube)", "https://www.youtube.com/@statquest")]),
    dict(id="genai", name="Generative AI & Agentic AI", short="GenAI & Agents", color="#EA580C", desc="How LLMs work, prompting, LLM APIs in Python, embeddings and vector search, RAG, tool use, agents, MCP, evaluation, guardrails, fine-tuning and responsible AI.",
         practice=[("W3Schools Generative AI tutorial", f"{W3}/gen_ai/"), ("Anthropic docs (Claude API)", "https://docs.claude.com/"), ("OpenAI cookbook", "https://cookbook.openai.com/"),
                   ("Hugging Face LLM course (free)", "https://huggingface.co/learn/llm-course"), ("Model Context Protocol docs", "https://modelcontextprotocol.io/"), ("DeepLearning.AI short courses", "https://www.deeplearning.ai/short-courses/")]),
    dict(id="career", name="Capstone, Portfolio & Interviews", short="Career", color="#475569", desc="End-to-end capstone, portfolio and GitHub, analyst and data scientist interview preparation with model answers, and a 90-day plan.",
         practice=[("W3Schools quizzes hub", f"{W3}/quiztest/"), ("W3Schools exercises hub", f"{W3}/exercises/"), ("LeetCode SQL 50", "https://leetcode.com/studyplan/top-sql-50/"), ("Kaggle competitions (Getting started)", "https://www.kaggle.com/competitions?hostSegmentIdFilter=5"),
                   ("StrataScratch data interview questions", "https://www.stratascratch.com/")]),
]
TRACK = {t["id"]: t for t in TRACKS}
BUILD_VER = "v14"

# ====================================================================== code highlighting
KW = {
    "python": r"def|return|if|elif|else|for|while|in|not|and|or|import|from|as|with|class|try|except|finally|lambda|True|False|None|break|continue|pass|is|yield|raise|global|async|await|match|case",
    "sql": r"SELECT|FROM|WHERE|GROUP|BY|ORDER|HAVING|JOIN|LEFT|RIGHT|FULL|INNER|OUTER|CROSS|APPLY|ON|AS|AND|OR|NOT|IN|IS|NULL|LIKE|BETWEEN|CASE|WHEN|THEN|ELSE|END|DISTINCT|TOP|UNION|ALL|INTERSECT|EXCEPT|WITH|OVER|PARTITION|ROWS|RANGE|PRECEDING|FOLLOWING|CURRENT|ROW|UNBOUNDED|CREATE|TABLE|VIEW|PROCEDURE|PROC|FUNCTION|ALTER|DROP|TRUNCATE|INSERT|INTO|VALUES|UPDATE|SET|DELETE|PRIMARY|KEY|FOREIGN|REFERENCES|DEFAULT|CHECK|UNIQUE|CONSTRAINT|IDENTITY|DECLARE|BEGIN|EXEC|EXECUTE|USE|DATABASE|ADD|COLUMN|OFFSET|FETCH|NEXT|FIRST|ONLY|EXISTS|ASC|DESC|LIMIT|PIVOT|UNPIVOT|FOR|INT|INTEGER|VARCHAR|NVARCHAR|CHAR|DATE|DATETIME|DECIMAL|FLOAT|BIT|MONEY|BIGINT|USING|ROLLUP|CUBE|GRANT|REVOKE|COMMIT|ROLLBACK|TRANSACTION|TRAN|SAVE|INDEX|CLUSTERED|NONCLUSTERED|IF|RETURN|RETURNS|GO|PRINT|OUTPUT|TRY|CATCH|RECURSIVE|MERGE|MATCHED|TRIGGER|AFTER|SCHEMA",
    "dax": r"VAR|RETURN|TRUE|FALSE|IN|NOT|AND|OR|DEFINE|EVALUATE|MEASURE|ORDER|BY",
    "m": r"let|in|each|type|true|false|null|if|then|else",
    "bash": r"git|pip|conda|docker|cd|python|uv|streamlit|uvicorn|npm|npx",
}


def highlight(code, lang):
    lang = (lang or "").lower()
    if lang in ("text", "", "output", "json", "yaml"):
        return html.escape(code)
    if lang == "excel":
        pat = re.compile(r'(?P<s>"[^"\n]*")|(?P<f>\b[A-Z][A-Z0-9.]+(?=\())|(?P<n>\b\d+(?:\.\d+)?\b)')
    else:
        com = {"python": r"#[^\n]*", "sql": r"--[^\n]*|/\*[\s\S]*?\*/", "dax": r"--[^\n]*|//[^\n]*", "m": r"//[^\n]*", "bash": r"#[^\n]*"}.get(lang, r"#[^\n]*")
        flags = re.I if lang in ("sql", "dax") else 0
        pat = re.compile(rf"(?P<c>{com})|(?P<s>'''[\s\S]*?'''|\"\"\"[\s\S]*?\"\"\"|f?\"[^\"\n]*\"|f?'[^'\n]*')|(?P<k>\b(?:{KW.get(lang, 'zzzz')})\b)|(?P<f>\b[A-Za-z_][A-Za-z0-9_.]*(?=\())|(?P<n>\b\d+(?:\.\d+)?\b)", flags)
    out, pos = [], 0
    for m in pat.finditer(code):
        out.append(html.escape(code[pos:m.start()]))
        out.append(f'<span class="hl-{m.lastgroup}">{html.escape(m.group())}</span>')
        pos = m.end()
    out.append(html.escape(code[pos:]))
    return "".join(out)


def codeblock(code, lang, label=None):
    lab = label or {"python": "Python", "sql": "SQL", "dax": "DAX", "m": "Power Query M", "excel": "Excel formula", "bash": "Terminal", "text": "Text", "json": "JSON", "yaml": "YAML"}.get(lang, lang.upper() if lang else "")
    return f'<div class="code"><div class="code-h"><span>{html.escape(lab)}</span><button class="copy" type="button">Copy</button></div><pre><code class="{lang}">{highlight(code, lang)}</code></pre></div>'


# ====================================================================== output rendering
def fmt_cell(x):
    if isinstance(x, (float, np.floating)):
        if np.isnan(x): return "NaN"
        if float(x).is_integer() and abs(x) < 1e15: return f"{x:,.0f}"
        return f"{x:,.2f}" if abs(x) >= .01 else f"{x:.4g}"
    if isinstance(x, (pd.Timestamp, datetime.datetime)):
        return x.strftime("%Y-%m-%d") if (x.hour, x.minute, x.second) == (0, 0, 0) else x.strftime("%Y-%m-%d %H:%M")
    if x is None or (isinstance(x, float) and np.isnan(x)): return "NULL"
    return x


def df_html(o, maxrows=12, index=True, na="NaN"):
    if isinstance(o, pd.Series):
        o = o.to_frame()
    n, m = o.shape
    d = (o.head(maxrows) if n > maxrows else o).copy()
    for c in d.columns:
        col = d[c]
        if pd.api.types.is_float_dtype(col):
            v = col.dropna().astype(float)
            if len(v) and (v.abs() < 0.01).all() and (v != 0).any():
                d[c] = col.map(lambda x: na if pd.isna(x) else f"{x:.4g}")
                continue
            dec = 2
            for k in (0, 1, 2):
                if len(v) == 0 or np.allclose(v, v.round(k), atol=1e-9): dec = k; break
            if dec == 0 and len(v) and v.abs().max() < 1e4 and (v != v.round(0)).any(): dec = 1
            d[c] = col.map(lambda x, dd=dec: na if pd.isna(x) else f"{x:,.{dd}f}")
        else:
            d[c] = col.map(lambda x: na if (x is None or x is pd.NA or x is pd.NaT or (isinstance(x, float) and np.isnan(x))) else fmt_cell(x))
    h = d.to_html(classes="tbl", border=0, escape=True, index=index, na_rep="NaN")
    foot = f'<div class="shape">{n:,} rows × {m} columns{" (first %d shown)" % maxrows if n > maxrows else ""}</div>'
    return f'<div class="tblwrap">{h}</div>{foot}'


# ====================================================================== python runner
class PyRunner:
    def __init__(self, lid):
        self.lid = lid; self.n = 0
        self.ns = {"__name__": "__main__", "dg": dg}
        self.errors = []

    def run(self, code, fig_only=False, caption=None, hide_code=False, expect_error=False):
        cwd = os.getcwd(); os.chdir(DATA)
        buf = io.StringIO(); err = None; result = None
        plt.close("all"); dg.style()
        try:
            tree = ast.parse(code)
            last = tree.body[-1] if tree.body and isinstance(tree.body[-1], ast.Expr) else None
            body = tree.body[:-1] if last is not None else tree.body
            with contextlib.redirect_stdout(buf):
                exec(compile(ast.Module(body=body, type_ignores=[]), f"<{self.lid}>", "exec"), self.ns)
                if last is not None:
                    result = eval(compile(ast.Expression(body=last.value), f"<{self.lid}>", "eval"), self.ns)
        except Exception as e:
            err = "".join(traceback.format_exception_only(type(e), e)).strip()
            if not expect_error: self.errors.append(f"{self.lid}: {err}\n---\n{code[:400]}")
        finally:
            os.chdir(cwd)
        imgs = []
        figs = [plt.figure(n) for n in plt.get_fignums()]
        if isinstance(result, plt.Figure) and result not in figs:
            figs.append(result)
        if isinstance(result, (plt.Figure,)) or type(result).__name__ in ("Axes", "AxesSubplot", "list", "Text") and figs:
            result = None if isinstance(result, plt.Figure) or type(result).__name__ in ("Axes", "AxesSubplot", "Text") else result
        if isinstance(result, list) and result and all(type(r).__name__ in ("Line2D", "Text", "Rectangle", "BarContainer") for r in result):
            result = None
        type_nm = type(result).__name__
        if type_nm in ("BarContainer", "Legend", "PathCollection", "tuple") and figs and type_nm != "tuple":
            result = None
        for f in figs:
            if not f.axes: continue
            self.n += 1
            name = f"{self.lid}-{self.n:02d}.png"
            try: f.tight_layout()
            except Exception: pass
            f.savefig(os.path.join(IMG, name), bbox_inches="tight", facecolor="white")
            imgs.append(name)
        plt.close("all")
        parts = []
        if fig_only:
            for im in imgs:
                parts.append(f'<figure class="fig"><img src="../img/{im}" alt="{html.escape(caption or "")}" loading="lazy">' + (f"<figcaption>{caption}</figcaption>" if caption else "") + "</figure>")
            if err: parts.append(f'<pre class="err">{html.escape(err)}</pre>')
            return "".join(parts)
        if not hide_code:
            parts.append(codeblock(code.strip("\n"), "python"))
        outs = []
        so = buf.getvalue().rstrip()
        if so: outs.append(f'<pre class="stdout">{html.escape(so)}</pre>')
        if result is not None and not (isinstance(result, (pd.DataFrame, pd.Series)) is False and type_nm in ("Axes",)):
            if isinstance(result, (pd.DataFrame, pd.Series)):
                outs.append(df_html(result))
            elif type(result).__name__ == "Styler":
                outs.append(df_html(result.data))
            else:
                outs.append(f'<pre class="stdout">{html.escape(repr(result))}</pre>')
        for im in imgs:
            outs.append(f'<figure class="fig"><img src="../img/{im}" alt="Output chart" loading="lazy">' + (f"<figcaption>{caption}</figcaption>" if caption else "") + "</figure>")
        if err: outs.append(f'<pre class="err">{html.escape(err)}</pre>')
        if outs:
            parts.append('<div class="output"><div class="out-h">Output</div>' + "".join(outs) + "</div>")
        return "".join(parts)


# ====================================================================== SQL runner (T-SQL shown, DuckDB executes)
def tsql_to_duck(q):
    s = q
    s = re.sub(r"\bGO\b", "", s)
    s = re.sub(r"\[([^\]]+)\]", r'"\1"', s)
    # TOP n  -> LIMIT n
    m = re.search(r"\bSELECT\s+(DISTINCT\s+)?TOP\s*\(?\s*(\d+)\s*\)?\s*(PERCENT\s+)?", s, re.I)
    if m and not m.group(3):
        s = s[:m.start()] + "SELECT " + (m.group(1) or "") + s[m.end():]
        s = s.rstrip().rstrip(";") + f"\nLIMIT {m.group(2)};"
    s = re.sub(r"OFFSET\s+(\d+)\s+ROWS?\s+FETCH\s+(?:NEXT|FIRST)\s+(\d+)\s+ROWS?\s+ONLY", r"LIMIT \2 OFFSET \1", s, flags=re.I)
    s = re.sub(r"OFFSET\s+(\d+)\s+ROWS?(?!\s+FETCH)", r"OFFSET \1", s, flags=re.I)
    s = re.sub(r"\bGETDATE\s*\(\s*\)", "current_timestamp", s, flags=re.I)
    s = re.sub(r"\bDATEPART\s*\(\s*(\w+)\s*,", lambda m: f"date_part('{ {'yy': 'year', 'yyyy': 'year', 'mm': 'month', 'dd': 'day', 'qq': 'quarter', 'dw': 'dow', 'weekday': 'dow', 'wk': 'week'}.get(m.group(1).lower(), m.group(1).lower()) }',", s, flags=re.I)
    s = re.sub(r"\bDATENAME\s*\(\s*month\s*,", "monthname(", s, flags=re.I)
    s = re.sub(r"\bDATENAME\s*\(\s*(?:weekday|dw)\s*,", "dayname(", s, flags=re.I)
    s = re.sub(r"\bDATEDIFF\s*\(\s*(\w+)\s*,", lambda m: f"date_diff('{m.group(1).lower()}',", s, flags=re.I)
    s = re.sub(r"\bDATEADD\s*\(\s*(\w+)\s*,\s*([^,]+?)\s*,", lambda m: f"date_add_sn('{m.group(1).lower()}', {m.group(2)},", s, flags=re.I)
    s = re.sub(r"\bEOMONTH\s*\(", "last_day(", s, flags=re.I)
    s = re.sub(r"\bISNULL\s*\(", "coalesce(", s, flags=re.I)
    s = re.sub(r"\bLEN\s*\(", "length(", s, flags=re.I)
    s = re.sub(r"\bCHARINDEX\s*\(\s*([^,]+?)\s*,\s*([^)]+?)\s*\)", r"instr(\2, \1)", s, flags=re.I)
    s = re.sub(r"\bNVARCHAR\b", "VARCHAR", s, flags=re.I)
    s = re.sub(r"\bVARCHAR\s*\(\s*MAX\s*\)", "VARCHAR", s, flags=re.I)
    s = re.sub(r"\bDATETIME2?\b", "TIMESTAMP", s, flags=re.I)
    s = re.sub(r"\bMONEY\b", "DECIMAL(19,4)", s, flags=re.I)
    s = re.sub(r"\bBIT\b", "BOOLEAN", s, flags=re.I)
    s = re.sub(r"\bSTDEV\s*\(", "stddev_samp(", s, flags=re.I)
    s = re.sub(r"\bVAR\s*\(", "var_samp(", s, flags=re.I)
    s = re.sub(r"\bCEILING\s*\(", "ceil(", s, flags=re.I)
    s = re.sub(r"\bSTRING_AGG\s*\(", "string_agg(", s, flags=re.I)
    s = re.sub(r"\bPERCENTILE_(CONT|DISC)\s*\(\s*([\d.]+)\s*\)\s*WITHIN\s+GROUP\s*\(\s*ORDER\s+BY\s+([^)]+?)\s*\)", lambda m: f"quantile_{m.group(1).lower()}({m.group(3)}, {m.group(2)})", s, flags=re.I)
    s = re.sub(r"\bWITHIN\s+GROUP\s*\(\s*ORDER\s+BY[^)]*\)", "", s, flags=re.I)
    s = re.sub(r"\bINT\s+IDENTITY\s*\(\s*\d+\s*,\s*\d+\s*\)", "INTEGER", s, flags=re.I)
    s = re.sub(r"\bDATEFROMPARTS\s*\(", "make_date(", s, flags=re.I)
    s = re.sub(r"\bCONVERT\s*\(\s*DATE\s*,\s*([^)]+)\)", r"CAST(\1 AS DATE)", s, flags=re.I)
    s = re.sub(r"\bIIF\s*\(", "if(", s, flags=re.I)
    s = re.sub(r"\bSQUARE\s*\(", "pow2(", s, flags=re.I)
    s = re.sub(r"\bREPLICATE\s*\(", "repeat(", s, flags=re.I)
    m = re.match(r"\s*WITH\s+(\w+)\s*(?:\([^)]*\))?\s+AS\s*\(", s, re.I)
    if m and re.search(rf"\b(?:FROM|JOIN)\s+{m.group(1)}\b", s, re.I) and re.search(r"UNION\s+ALL", s, re.I) and not re.match(r"\s*WITH\s+RECURSIVE", s, re.I):
        s = re.sub(r"^\s*WITH\b", "WITH RECURSIVE", s, count=1, flags=re.I)
    s = re.sub(r"\bOPTION\s*\(\s*MAXRECURSION\s+\d+\s*\)", "", s, flags=re.I)
    s = re.sub(r"\bPRINT\b[^\n]*", "", s, flags=re.I)
    s = re.sub(r"\bUSE\s+\w+\s*;?", "", s, flags=re.I)
    return s


class SQLRunner:
    def __init__(self, lid):
        import duckdb
        self.lid = lid; self.errors = []
        self.con = duckdb.connect()
        for nm, f in [("superstore", "superstore_sales.csv"), ("customers", "customers.csv"), ("products", "products.csv"), ("employees", "hr_employees.csv"),
                      ("video_games", "video_game_sales.csv"), ("emp", "employees_small.csv"), ("houses", "house_prices.csv"), ("daily_sales", "daily_sales.csv")]:
            self.con.execute(f"CREATE TABLE {nm} AS SELECT * FROM read_csv_auto('{os.path.join(DATA, f)}')")
        self.con.execute("CREATE TABLE orders AS SELECT DISTINCT order_id, order_date, ship_date, ship_mode, customer_id FROM superstore ORDER BY order_date, order_id")
        self.con.execute("CREATE TABLE order_items AS SELECT row_id AS item_id, order_id, product_id, quantity, discount, sales, profit FROM superstore ORDER BY row_id")
        self.con.execute("CREATE MACRO date_add_sn(part, n, d) AS CASE part WHEN 'day' THEN d + to_days(CAST(n AS INTEGER)) WHEN 'dd' THEN d + to_days(CAST(n AS INTEGER)) WHEN 'month' THEN d + to_months(CAST(n AS INTEGER)) WHEN 'mm' THEN d + to_months(CAST(n AS INTEGER)) WHEN 'year' THEN d + to_years(CAST(n AS INTEGER)) WHEN 'week' THEN d + to_days(7*CAST(n AS INTEGER)) ELSE d END")
        self.con.execute("CREATE MACRO pow2(x) AS x*x")

    def run(self, code, show_code=True, maxrows=12):
        parts = [codeblock(code.strip("\n"), "sql")] if show_code else []
        stmts = [s for s in split_sql(code) if s.strip()]
        res = None; err = None; msg = []; results = []
        for st in stmts:
            q = tsql_to_duck(st)
            if not q.strip().strip(";").strip(): continue
            try:
                kind = re.sub(r"--[^\n]*", "", st).strip().split()[0].upper()
                r = self.con.execute(q)
                if kind in ("INSERT", "UPDATE", "DELETE", "MERGE"):
                    n = None
                    try:
                        d = r.df(); n = int(d.iloc[0, 0]) if len(d) else 0
                    except Exception: pass
                    msg.append(f"({n} row{'s' if n != 1 else ''} affected)" if n is not None else "Rows affected.")
                elif kind in ("BEGIN", "COMMIT", "ROLLBACK", "SAVE", "SET"):
                    msg.append("Commands completed successfully.")
                elif kind in ("CREATE", "ALTER", "DROP", "TRUNCATE"):
                    msg.append("Commands completed successfully.")
                elif r.description:
                    res = r.df(); results.append(res)
            except Exception as e:
                err = str(e).split("\n")[0]
                self.errors.append(f"{self.lid}: {err}\n---\n{q[:500]}")
                break
        outs = []
        for k, res_k in enumerate(results):
            if "count_star()" in res_k.columns: res_k = res_k.rename(columns={"count_star()": "(No column name)"})
            outs.append((f'<div class="rs-n">Result set {k + 1}</div>' if len(results) > 1 else "") + df_html(res_k, maxrows=maxrows, index=False, na="NULL"))
        if msg and not err and (not results or any("affected" in m for m in msg)):
            outs.insert(0, f'<pre class="stdout msgs">{html.escape(chr(10).join(msg))}</pre>')
        if err:
            outs.append(f'<pre class="err">{html.escape(err)}</pre>')
        if outs:
            parts.append('<div class="output"><div class="out-h">Result</div>' + "".join(outs) + "</div>")
        return "".join(parts)


def split_sql(code):
    out, cur, q = [], [], False
    for ch in code:
        if ch == "'": q = not q
        if ch == ";" and not q:
            out.append("".join(cur)); cur = []
        else:
            cur.append(ch)
    out.append("".join(cur))
    return [o for o in out if re.sub(r"--[^\n]*", "", o).strip()]


# ====================================================================== Excel grids (LibreOffice recalculation)
XLFN = ["XLOOKUP", "XMATCH", "IFS", "SWITCH", "MAXIFS", "MINIFS", "CONCAT", "TEXTJOIN", "UNIQUE", "SORTBY", "SEQUENCE", "LET", "LAMBDA", "CEILING.MATH", "FLOOR.MATH",
        "IFNA", "DAYS", "ISOWEEKNUM", "STDEV.S", "STDEV.P", "VAR.S", "VAR.P", "PERCENTILE.INC", "QUARTILE.INC", "RANK.EQ", "AGGREGATE", 
        "CHOOSECOLS", "TAKE", "DROP", "VSTACK", "HSTACK", "TEXTSPLIT", "TEXTBEFORE", "TEXTAFTER", "RANDARRAY", "NORM.DIST", "T.TEST", "MODE.SNGL", "FORECAST.LINEAR", "CONFIDENCE.NORM", "GROUPBY", "PIVOTBY"]


def xl_prefix(f):
    for fn in sorted(XLFN, key=len, reverse=True):
        f = re.sub(rf"(?<![A-Z._]){re.escape(fn)}\(", f"_xlfn.{fn}(", f)
    f = re.sub(r"(?<![A-Z._])(SORT|FILTER)\(", r"_xlfn._xlws.\1(", f)
    return f


def col_letter(i):
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26); s = chr(65 + r) + s
    return s


def parse_cell(raw):
    raw = raw.strip()
    cls, fmt, override = "", None, None
    m = re.search(r"\{([a-z!]*)(?:;([^}]*))?\}\s*$", raw)
    if m:
        cls, fmt = m.group(1), m.group(2); raw = raw[:m.start()].rstrip()
    if len(raw) > 1 and raw.startswith('"') and raw.endswith('"') and raw.count('"') == 2:
        return dict(v=raw[1:-1], f=None, cls=cls, fmt=fmt, ov=None)
    if raw.startswith("=") and "=>>" in raw:
        raw, override = [x.strip() for x in raw.split("=>>", 1)]
        if override.startswith("!"): override = override[1:].strip(); cls = cls + "!"
    if raw.startswith("'"):
        return dict(v=raw[1:], f=None, cls=cls, fmt=fmt, ov=None)
    if raw.startswith("="):
        return dict(v=None, f=raw, cls=cls, fmt=fmt, ov=override)
    v = raw
    if re.fullmatch(r"-?\d+", raw): v = int(raw)
    elif re.fullmatch(r"-?\d*\.\d+", raw): v = float(raw)
    elif re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw): v = datetime.date.fromisoformat(raw)
    elif raw in ("TRUE", "FALSE"): v = raw == "TRUE"
    elif raw.endswith("%") and re.fullmatch(r"-?\d+(\.\d+)?%", raw): v = float(raw[:-1]) / 100; fmt = fmt or ("0%" if "." not in raw else "0.0%")
    return dict(v=v, f=None, cls=cls, fmt=fmt, ov=None)


def parse_xl(body):
    rows = []
    for line in body.strip("\n").split("\n"):
        if not line.strip(): rows.append([]); continue
        rows.append([parse_cell(c) for c in split_cells(line)])
    return rows


def split_cells(line):
    out, cur, q = [], [], False
    for ch in line:
        if ch == '"': q = not q
        if ch == "|" and not q: out.append("".join(cur)); cur = []
        else: cur.append(ch)
    out.append("".join(cur))
    return out


XLCACHE_F = os.path.join(CACHE, "xl.json")
XLCACHE = json.load(open(XLCACHE_F)) if os.path.exists(XLCACHE_F) else {}


def xl_compute(blocks):
    """blocks: {hash: rows}. Fills XLCACHE[hash] = {cellref: value}."""
    todo = {h: r for h, r in blocks.items() if h not in XLCACHE}
    if not todo: return
    import openpyxl
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    names = {}
    for h, (rows, ws_) in todo.items():
        for nm, csv in ws_: add_data_sheet(wb, nm, csv)
    for i, (h, (rows, ws_)) in enumerate(todo.items()):
        ws = wb.create_sheet(f"S{i}"); names[h] = f"S{i}"
        for r, row in enumerate(rows, 1):
            for c, cell in enumerate(row, 1):
                if cell["f"]: ws.cell(r, c, xl_prefix(cell["f"]))
                elif cell["v"] not in ("", None): ws.cell(r, c, cell["v"])
    tmp = os.path.join(CACHE, "xlcalc"); shutil.rmtree(tmp, ignore_errors=True); os.makedirs(tmp)
    wb.save(os.path.join(tmp, "in.xlsx"))
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", os.path.join(tmp, "out"), os.path.join(tmp, "in.xlsx")], capture_output=True, timeout=600)
    wb2 = openpyxl.load_workbook(os.path.join(tmp, "out", "in.xlsx"), data_only=True)
    for h, sn in names.items():
        ws = wb2[sn]; vals = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None: continue
                if isinstance(v, datetime.datetime): v = {"__date__": v.isoformat()}
                elif isinstance(v, datetime.time): v = {"__time__": (v.hour * 3600 + v.minute * 60 + v.second) / 86400}
                elif isinstance(v, datetime.timedelta): v = {"__time__": v.total_seconds() / 86400}
                vals[c.coordinate] = v
        XLCACHE[h] = vals
    json.dump(XLCACHE, open(XLCACHE_F, "w"))


def _serial(d):
    return (d - datetime.datetime(1899, 12, 30)).total_seconds() / 86400


def fmt_xl(v, fmt):
    if isinstance(v, dict) and "__date__" in v and fmt and re.fullmatch(r"[0.#,]+", fmt):
        x = _serial(datetime.datetime.fromisoformat(v["__date__"]))
        dec = len(fmt.split(".")[1]) if "." in fmt else 0
        return f"{x:.{dec}f}"
    if isinstance(v, (int, float)) and not isinstance(v, bool) and fmt and re.search(r"yy|mmm|dd-", fmt):
        v = {"__date__": (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(v))).isoformat()}
    if isinstance(v, dict) and "__time__" in v:
        f = v["__time__"]
        if fmt and re.fullmatch(r"[0.]+", fmt): return f"{f:.{len(fmt.split('.')[1]) if '.' in fmt else 0}f}"
        t = round(f * 1440); return f"{t // 60:02d}:{t % 60:02d}"
    if isinstance(v, dict) and "__date__" in v:
        d = datetime.datetime.fromisoformat(v["__date__"])
        if fmt and "h" in fmt and ":" in fmt: return d.strftime("%d-%m-%Y %H:%M") if "y" in fmt else d.strftime("%H:%M")
        if fmt == "mmm-yy": return d.strftime("%b-%y")
        if fmt and "mmm" in fmt: return d.strftime("%d-%b-%Y")
        return d.strftime("%d-%m-%Y")
    if isinstance(v, datetime.date): return v.strftime("%d-%m-%Y") if not fmt or "mmm" not in fmt else v.strftime("%d-%b-%Y")
    if isinstance(v, bool): return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        if fmt:
            if fmt.endswith("%"):
                dec = len(fmt.split(".")[1]) - 1 if "." in fmt else 0
                return f"{v * 100:.{dec}f}%"
            cur = "₹" if "₹" in fmt else ("$" if "$" in fmt else "")
            dec = len(fmt.split(".")[1]) if "." in fmt else 0
            if "#,##" in fmt: return f"{cur}{v:,.{dec}f}"
            if re.fullmatch(r"[₹$]?0(\.0+)?", fmt): return f"{cur}{v:.{dec}f}"
            if fmt == "hh:mm":
                t = (v % 1) * 1440; return f"{int(t // 60):02d}:{int(round(t % 60)):02d}"
        if isinstance(v, float) and not float(v).is_integer():
            if abs(v) >= 1e6: return f"{v:,.0f}"
            if abs(v) >= 100: return f"{v:,.2f}".rstrip("0").rstrip(".")
            return f"{v:,.4f}".rstrip("0").rstrip(".")
        if isinstance(v, float): v = int(v)
        return f"{v}"
    return html.escape(str(v))


def render_xl(rows, opts, h):
    vals = XLCACHE.get(h, {})
    show_f = "formulas" in opts
    title = opts.get("title")
    ncol = max((len(r) for r in rows), default=1)
    # spills from overrides
    spill = {}
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            if cell["f"] and cell["ov"] is not None and ";" in cell["ov"]:
                for k, part in enumerate(cell["ov"].split(";")):
                    for j, sub in enumerate(part.split("~")):
                        spill[(r + k, c + j)] = sub.strip()
                        ncol = max(ncol, c + j + 1)
    nrow = max(len(rows), max((k[0] + 1 for k in spill), default=0))
    th = '<tr><th class="corner"></th>' + "".join(f"<th>{col_letter(c)}</th>" for c in range(ncol)) + "</tr>"
    body = []
    flist = []
    issues = []
    for r in range(nrow):
        row = rows[r] if r < len(rows) else []
        tds = []
        for c in range(ncol):
            cell = row[c] if c < len(row) else None
            ref = f"{col_letter(c)}{r + 1}"
            classes = []
            if cell is None or (cell["f"] is None and cell["v"] in ("", None)):
                txt = html.escape(spill.get((r, c), "")) if (r, c) in spill else ""
                if (r, c) in spill: classes.append("spill")
                tds.append(f'<td class="{" ".join(classes)}">{txt}</td>')
                continue
            for ch in cell["cls"]:
                classes.append({"g": "c-g", "r": "c-r", "y": "c-y", "b": "c-b", "h": "c-h", "p": "c-p", "i": "c-i", "o": "c-o", "s": "c-s"}.get(ch, ""))
            title_attr = ""
            if cell["f"]:
                classes.append("fx")
                flist.append((ref, cell["f"]))
                if show_f:
                    txt = html.escape(cell["f"])
                else:
                    if cell["ov"] is not None:
                        ovv = cell["ov"].split(";")[0].split("~")[0].strip()
                        txt = html.escape(ovv)
                        got = vals.get(ref)
                        if got is not None and "!" not in cell["cls"] and not (isinstance(got, str) and got.startswith("#")) and ";" not in cell["ov"]:
                            if html.unescape(fmt_xl(got, cell["fmt"])) != ovv: issues.append(f"{ref}: LibreOffice={fmt_xl(got, cell['fmt'])} override={ovv}")
                    else:
                        got = vals.get(ref)
                        txt = fmt_xl(got, cell["fmt"]) if got is not None else ""
                        if isinstance(got, str) and got.startswith("#") and not opts.get("errors_ok"):
                            issues.append(f"{ref}: {got} for {cell['f']}")
                title_attr = f' title="{html.escape(cell["f"])}"'
            else:
                v = cell["v"]
                txt = fmt_xl(v, cell["fmt"]) if not isinstance(v, str) else html.escape(v)
                if isinstance(v, (int, float)) and not isinstance(v, bool): classes.append("num")
            if cell["f"] and not show_f:
                got = vals.get(ref)
                if isinstance(got, (int, float)) and not isinstance(got, bool) or (cell["ov"] and re.fullmatch(r"[-₹$\d.,%]+", cell["ov"].split(";")[0])): classes.append("num")
            tds.append(f'<td class="{" ".join(x for x in classes if x)}"{title_attr}>{txt}</td>')
        body.append(f'<tr><th>{r + 1}</th>{"".join(tds)}</tr>')
    cap = f'<div class="xl-title"><span class="xl-ico">X</span>{html.escape(title)}</div>' if title else ""
    legend = ""
    if flist and not show_f:
        seen, items = set(), []
        for ref, f in flist:
            key = re.sub(r"\d+", "#", f)
            if key in seen and len(flist) > 6: continue
            seen.add(key); items.append(f'<li><b>{ref}</b><code>{highlight(f, "excel")}</code></li>')
        more = '<li class="muted">Each pattern is copied down/across to the other cells in its column.</li>' if len(items) < len(flist) else ""
        legend = f'<ul class="xl-f">{"".join(items[:14])}{more}</ul>'
    view = '<span class="xl-mode">Formula view (Ctrl + `)</span>' if show_f else ""
    return f'<div class="xl">{cap}{view}<div class="xl-scroll"><table class="sheet">{th}{"".join(body)}</table></div>{legend}</div>', issues


# ====================================================================== markdown pipeline
CALLOUT = {"note": "Note", "tip": "Tip", "warn": "Watch out", "example": "Example", "live": "Live class", "new": "2026 update", "interview": "Interview corner",
           "key": "Key takeaways", "mistake": "Common mistakes", "ai": "Do it with AI", "practice": "Practice", "case": "Case study", "def": "Definition", "steps": "Step by step"}
MD_EXT = ["tables", "attr_list", "md_in_html", "sane_lists", "toc", "def_list"]


def md(text):
    return markdown.markdown(text, extensions=MD_EXT, extension_configs={"toc": {"permalink": False}})


def parse_opts(info):
    opts = {}
    for m in re.finditer(r'(\w+)="([^"]*)"|(\w+)=(\S+)|(\w+)', info):
        if m.group(1): opts[m.group(1)] = m.group(2)
        elif m.group(3): opts[m.group(3)] = m.group(4)
        else: opts[m.group(5)] = True
    return opts


def render_quiz(body, qid):
    qs = []
    cur = None
    for line in body.strip().split("\n"):
        s = line.strip()
        if not s: continue
        if s.startswith("Q:"):
            cur = dict(q=s[2:].strip(), opts=[], ans=None, exp=""); qs.append(cur)
        elif s.startswith("- ") or s.startswith("* "):
            if s.startswith("* "): cur["ans"] = len(cur["opts"])
            cur["opts"].append(s[2:].strip())
        elif s.startswith(">"):
            cur["exp"] += " " + s[1:].strip()
    out = [f'<div class="quiz" id="{qid}"><div class="quiz-h"><b>Quick quiz</b><span>{len(qs)} questions · tap an answer to check it</span></div>']
    for i, q in enumerate(qs):
        opts = "".join(f'<button type="button" class="opt" data-ok="{1 if j == q["ans"] else 0}"><i>{chr(65 + j)}</i>{inline_md(o)}</button>' for j, o in enumerate(q["opts"]))
        out.append(f'<div class="qq"><p class="qtext"><b>{i + 1}.</b> {inline_md(q["q"])}</p><div class="opts">{opts}</div>'
                   f'<div class="qexp"><b>Answer: {chr(65 + q["ans"])}.</b> {inline_md(q["exp"].strip())}</div></div>')
    out.append("</div>")
    return "".join(out), qs


def inline_md(s):
    h = md(s)
    return re.sub(r"^<p>|</p>$", "", h.strip())


class Lesson:
    def __init__(self, path):
        raw = open(path, encoding="utf-8").read()
        m = re.match(r"---\n([\s\S]*?)\n---\n", raw)
        fm = []
        for line in m.group(1).split("\n"):
            mm = re.match(r"^(\s*-\s+|\w+:\s+)(.*)$", line)
            if mm and ": " in mm.group(2) and not mm.group(2).lstrip().startswith(('"', "'", "[")):
                line = mm.group(1) + '"' + mm.group(2).replace('\\', '\\\\').replace('"', '\\"') + '"'
            fm.append(line)
        self.meta = yaml.safe_load("\n".join(fm)); self.body = raw[m.end():]
        fix = lambda o: ", ".join(f"{k}: {v}" for k, v in o.items()) if isinstance(o, dict) else str(o)
        self.meta["outcomes"] = [fix(o) for o in self.meta.get("outcomes", [])]
        for k in ("summary", "title", "live", "updated"):
            if isinstance(self.meta.get(k), dict): self.meta[k] = fix(self.meta[k])
        self.path = path; self.track = os.path.basename(os.path.dirname(path))
        self.num = int(self.meta["num"]); self.title = self.meta["title"]
        self.id = f"{self.track}-{self.num:02d}"
        self.slug = f"{self.id}-" + re.sub(r"[^a-z0-9]+", "-", self.title.lower().replace("&", "and")).strip("-")[:60].strip("-")
        self.file = self.slug + ".html"
        tn = TRACK[self.track]["short"].replace(" & ", "-").replace(" ", "-")
        self.pdf = f"{TRACK[self.track]['short'].replace(' & ', '_').replace(' ', '_')}_{self.num:02d}_" + re.sub(r"[^A-Za-z0-9]+", "_", self.title).strip("_")[:70] + ".pdf"
        self.hash = hashlib.sha1((raw + BUILD_VER + open(os.path.join(SRC, "dg.py")).read()).encode()).hexdigest()[:16]


def xl_hash(code, opts):
    return hashlib.sha1((code + "|" + str(opts.get("with", ""))).encode()).hexdigest()[:16]


def withs(opts):
    w = opts.get("with")
    return [tuple(x.split(":", 1)) for x in w.split(",")] if w else []


def collect_xl(lessons):
    blocks = {}
    for L in lessons:
        for m in re.finditer(r"^```xl([^\n]*)\n([\s\S]*?)^```", L.body, re.M):
            o = parse_opts(m.group(1).strip())
            blocks[xl_hash(m.group(2), o)] = (parse_xl(m.group(2)), withs(o))
    xl_compute(blocks)


def add_data_sheet(wb, name, csv):
    if name in wb.sheetnames: return
    ws = wb.create_sheet(name)
    df = pd.read_csv(os.path.join(DATA, csv))
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        vals = []
        for v in row:
            if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v): v = datetime.date.fromisoformat(v)
            elif isinstance(v, (np.integer,)): v = int(v)
            elif isinstance(v, (np.floating,)): v = None if np.isnan(v) else float(v)
            vals.append(v)
        ws.append(vals)


def render_lesson(L):
    cache_f = os.path.join(CACHE, L.id + ".json")
    if os.path.exists(cache_f):
        c = json.load(open(cache_f))
        if c.get("hash") == L.hash and all(os.path.exists(os.path.join(IMG, i)) for i in c.get("imgs", [])):
            return c
    for old in glob.glob(os.path.join(IMG, L.id + "-*.png")): os.remove(old)
    py = PyRunner(L.id); sq = None; issues = []
    blocks = []
    xl_blocks = []

    def take_block(m):
        info, code = m.group(1).strip(), m.group(2)
        lang = info.split()[0] if info else ""
        rest = info[len(lang):].strip()
        if lang == "python" and rest.startswith("fig"):
            cap = rest.split(":", 1)[1].strip() if ":" in rest else None
            h = py.run(code, fig_only=True, caption=inline_md(cap) if cap else None)
        elif lang == "python" and rest.startswith("run"):
            o = parse_opts(rest)
            h = py.run(code, hide_code="hide" in o, caption=o.get("caption"), expect_error="error" in o)
        elif lang == "python" and rest.startswith("setup"):
            py.run(code, hide_code=True); h = ""
        elif lang == "sql" and rest.startswith("run"):
            nonlocal sq
            if sq is None: sq = SQLRunner(L.id)
            o = parse_opts(rest)
            h = sq.run(code, maxrows=int(o.get("rows", 12)))
        elif lang == "xl":
            o = parse_opts(rest); rows = parse_xl(code); hh = xl_hash(code, o)
            h, iss = render_xl(rows, o, hh); issues.extend(iss); xl_blocks.append((o.get("title", f"Sheet{len(xl_blocks) + 1}"), rows, o))
        else:
            o = parse_opts(rest)
            h = codeblock(code.rstrip("\n"), lang, o.get("label"))
        blocks.append(h)
        return f"\n\nXQBLOCK{len(blocks) - 1}Z\n\n"

    body = re.sub(r"^```([^\n]*)\n([\s\S]*?)^```[ \t]*$", take_block, L.body, flags=re.M)
    quizzes = []

    def callout(m):
        kind, title, inner = m.group(1), (m.group(2) or "").strip(), m.group(3)
        if kind == "quiz":
            h, qs = render_quiz(inner, f"quiz-{L.id}-{len(quizzes)}"); quizzes.extend(qs)
            blocks.append(h); return f"\n\nXQBLOCK{len(blocks) - 1}Z\n\n"
        if kind == "exercise":
            prob, _, sol = inner.partition("???")
            h = (f'<div class="exercise"><div class="ex-h"><span class="pill">Exercise</span><b>{html.escape(title)}</b></div>{fill(md(prob))}'
                 + (f'<details class="sol"><summary>Show solution</summary>{fill(md(sol))}</details>' if sol.strip() else "") + "</div>")
            blocks.append(h); return f"\n\nXQBLOCK{len(blocks) - 1}Z\n\n"
        kind = {"warning": "warn"}.get(kind, kind)
        lab = CALLOUT.get(kind, kind.title())
        h = f'<div class="co co-{kind}"><div class="co-h"><span class="pill">{lab}</span>{"<b>" + html.escape(title) + "</b>" if title else ""}</div>{fill(md(inner))}</div>'
        blocks.append(h); return f"\n\nXQBLOCK{len(blocks) - 1}Z\n\n"

    def fill(h):
        for _ in range(3):
            h = re.sub(r"<p>XQBLOCK(\d+)Z</p>", lambda m: blocks[int(m.group(1))], h)
            h = re.sub(r"XQBLOCK(\d+)Z", lambda m: blocks[int(m.group(1))], h)
        return h

    body = re.sub(r"^:::(\w+)[ \t]*([^\n]*)\n([\s\S]*?)^:::[ \t]*$", callout, body, flags=re.M)
    h = fill(md(body))
    h = re.sub(r"<table>", '<div class="tblwrap"><table class="tbl md">', h).replace("</table>", "</table></div>")
    h = re.sub(r'<p><img alt="([^"]*)" src="([^"]+)"\s*/?></p>', lambda m: f'<figure class="fig"><img src="{m.group(2)}" alt="{m.group(1)}" loading="lazy"><figcaption>{m.group(1)}</figcaption></figure>', h)
    h = re.sub(r'<a href="(https?://[^"]+)"', r'<a target="_blank" rel="noopener" href="\1"', h)
    toc = [(m.group(1), re.sub(r"<[^>]+>", "", m.group(2))) for m in re.finditer(r'<h2 id="([^"]+)">(.*?)</h2>', h)]
    imgs = sorted(set(re.findall(r'\.\./img/([^"]+\.png)', h)))
    text = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", h))
    c = dict(hash=L.hash, html=h, toc=toc, imgs=imgs, errors=py.errors + (sq.errors if sq else []), issues=issues, quiz=quizzes,
             words=len(text.split()), xl=[(t, [[{k: (v.isoformat() if isinstance(v, (datetime.date,)) else v) for k, v in cell.items()} for cell in row] for row in rows], o) for t, rows, o in xl_blocks],
             search=text[:6000])
    json.dump(c, open(cache_f, "w"))
    return c


# ====================================================================== page templates
def esc(s): return html.escape(str(s))


HEAD = """<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title}</title><meta name="description" content="{desc}"><link rel="icon" href="{r}assets/logo.svg">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&family=JetBrains+Mono:wght@400;600&display=swap" rel="stylesheet">
<link rel="stylesheet" href="{r}assets/course.css?v={v}"></head><body class="{bodycls}">"""


def topbar(r, crumbs=""):
    return f"""<header class="top"><a class="brand" href="{r}index.html"><img src="{r}assets/logo.svg" alt="" width="30" height="30"><span><b>Skill<em>Navora</em></b><small>{TITLE}</small></span></a>
<nav class="crumbs">{crumbs}</nav><div class="top-r"><input id="q" type="search" placeholder="Search lessons…" aria-label="Search lessons"><div id="qres" class="qres" hidden></div>
<button id="theme" type="button" aria-label="Toggle dark mode">◐</button></div></header>"""


FOOT = """<footer class="foot"><div>© 2026 {brand} · {title} by {author} · Version {ver} ({upd})</div><div>Original teaching material. Practice datasets are fictional. Product names belong to their owners.</div></footer>
<script src="{r}assets/search-index.js?v={v}"></script><script src="{r}assets/course.js?v={v}"></script></body></html>"""


def lesson_page(L, c, prev, nxt, idx, n_in_track):
    t = TRACK[L.track]; m = L.meta
    crumbs = f'<a href="../index.html">Handbook</a><span>/</span><a href="../index.html#t-{t["id"]}">{esc(t["name"])}</a><span>/</span><b>Lesson {L.num}</b>'
    tags = []
    if m.get("live"): tags.append(f'<span class="tag live">● {esc(m["live"])}</span>')
    if m.get("updated"): tags.append(f'<span class="tag new">✦ {esc(m["updated"])}</span>')
    tags.append(f'<span class="tag">{esc(m.get("level", "Beginner"))}</span>')
    tags.append(f'<span class="tag">⏱ {esc(m.get("duration", "60 min"))}</span>')
    outcomes = "".join(f"<li>{inline_md(o)}</li>" for o in m.get("outcomes", []))
    toc = "".join(f'<a href="#{a}">{esc(b)}</a>' for a, b in c["toc"])
    nav = '<div class="pn">' + (f'<a class="prev" href="{prev.file}"><small>← Previous</small><b>{esc(prev.title)}</b></a>' if prev else "<span></span>") + \
          (f'<a class="next" href="{nxt.file}"><small>Next →</small><b>{esc(nxt.title)}</b></a>' if nxt else "<span></span>") + "</div>"
    practice = "".join(f'<li><a target="_blank" rel="noopener" href="{u}">{esc(n)} ↗</a></li>' for n, u in t["practice"])
    files = ""
    if os.path.exists(os.path.join(FILES, "excel", L.id + ".xlsx")):
        files = f'<a class="btn ghost" href="../files/excel/{L.id}.xlsx" download>⬇ Practice workbook (.xlsx)</a>'
    data = "".join(f'<a class="chip" href="../datasets/{d}" download>{esc(d)}</a>' for d in m.get("datasets", []))
    head = HEAD.format(title=esc(f"{L.title} · {t['short']} Lesson {L.num} · {TITLE}"), desc=esc(m.get("summary", t["desc"])), r="../", v=BUILD_VER, bodycls=f"lesson tr-{t['id']}")
    return head + topbar("../", crumbs) + f"""
<div class="layout"><aside class="side"><div class="side-in"><div class="side-t" style="--tc:{t['color']}">{esc(t['name'])}</div>
<ol class="side-l">{"".join(f'<li class="{"on" if x.id == L.id else ""}"><a href="{x.file}" data-lid="{x.id}"><span>{x.num:02d}</span>{esc(x.title)}</a></li>' for x in idx)}</ol></div></aside>
<main class="main" id="top"><article class="lesson-art" data-lid="{L.id}">
<div class="lhead" style="--tc:{t['color']}"><div class="kick">{esc(t['name'])} · Lesson {L.num} of {n_in_track}</div><h1>{esc(L.title)}</h1>
<p class="sum">{inline_md(m.get('summary', ''))}</p><div class="tags">{"".join(tags)}</div>
<div class="lacts"><a class="btn" href="../pdf/{L.track}/{L.pdf}" download>⬇ Download PDF</a>{files}<button class="btn ghost done" type="button" data-lid="{L.id}">☐ Mark as complete</button></div></div>
<div class="outcomes"><div><b>By the end of this lesson you can</b><ul>{outcomes}</ul></div>{f'<div class="dsets"><b>Datasets used</b><div>{data}</div></div>' if data else ''}</div>
<div class="content">{c['html']}</div>
<section class="practice"><h2 id="practice-more">Practise more online</h2><p>Free, trusted places to test yourself on this topic. Links open the original site in a new tab.</p><ul>{practice}</ul></section>
{nav}</article></main><nav class="toc"><div class="toc-in"><b>On this page</b>{toc}<a href="#practice-more">Practise more online</a></div></nav></div>
""" + FOOT.format(brand=BRAND, title=TITLE, author=AUTHOR, ver=VERSION, upd=UPDATED, r="../", v=BUILD_VER)


# ====================================================================== main
def main(only=None):
    files = sorted(glob.glob(os.path.join(SRC, "lessons", "*", "*.md")))
    lessons = [Lesson(f) for f in files]
    order = {t["id"]: i for i, t in enumerate(TRACKS)}
    lessons.sort(key=lambda L: (order[L.track], L.num))
    collect_xl(lessons)
    rendered = {}
    errs, iss = [], []
    for L in lessons:
        if only and not any(o in L.id for o in only):
            cf = os.path.join(CACHE, L.id + ".json")
            if os.path.exists(cf): rendered[L.id] = json.load(open(cf)); continue
        c = render_lesson(L); rendered[L.id] = c
        if c["errors"]: errs += c["errors"]
        if c["issues"]: iss += [f"{L.id} {x}" for x in c["issues"]]
        print(f"{L.id:12s} {c['words']:6d} words {len(c['imgs']):3d} imgs {len(c['quiz']):2d} quiz {'ERR ' + str(len(c['errors'])) if c['errors'] else ''}")
    # excel practice workbooks
    write_workbooks(lessons, rendered)
    by_track = {t["id"]: [L for L in lessons if L.track == t["id"]] for t in TRACKS}
    for t in TRACKS:
        idx = by_track[t["id"]]
        for i, L in enumerate(idx):
            prev = idx[i - 1] if i else None; nxt = idx[i + 1] if i + 1 < len(idx) else None
            open(os.path.join(LES, L.file), "w", encoding="utf-8").write(lesson_page(L, rendered[L.id], prev, nxt, idx, len(idx)))
    import hub
    hub.build(lessons, rendered, by_track, globals())
    if errs:
        print("\n==== EXECUTION ERRORS ====\n" + "\n\n".join(errs))
    if iss:
        print("\n==== EXCEL ISSUES ====\n" + "\n".join(iss))
    tot = sum(c["words"] for c in rendered.values())
    print(f"\n{len(lessons)} lessons, {tot:,} words, {sum(len(c['imgs']) for c in rendered.values())} images, {sum(len(c['quiz']) for c in rendered.values())} quiz questions")


def write_workbooks(lessons, rendered):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.formula import ArrayFormula
    os.makedirs(os.path.join(FILES, "excel"), exist_ok=True)
    fills = {"g": "C6EFCE", "r": "FFC7CE", "y": "FFEB9C", "h": "1F2A44", "p": "E6E0FF", "o": "FCE4D6", "s": "E2F7F4"}
    for L in lessons:
        xs = rendered[L.id].get("xl") or []
        if not xs: continue
        wb = openpyxl.Workbook(); wb.remove(wb.active)
        ws0 = wb.create_sheet("Read me")
        ws0["A1"] = f"{BRAND} · {TITLE}"; ws0["A1"].font = Font(bold=True, size=14, color="6C4CF5")
        ws0["A2"] = f"Lesson {L.num}: {L.title}"; ws0["A2"].font = Font(bold=True, size=12)
        ws0["A4"] = "Each sheet matches one example in the lesson. Formulas are live: change the inputs and watch the results update."
        ws0["A5"] = "Sample data is fictional. Created by Varun Kumar for learners."
        ws0.column_dimensions["A"].width = 100
        used = set()
        for title, rows, o in xs:
            for nm, csv in withs(o): add_data_sheet(wb, nm, csv); used.add(nm)
        for k, (title, rows, o) in enumerate(xs, 1):
            nm = re.sub(r"[\[\]:*?/\\]", "", f"{k}. {title}")[:31]
            while nm in used: nm = nm[:28] + str(k)
            used.add(nm); ws = wb.create_sheet(nm)
            for r, row in enumerate(rows, 1):
                for c, cell in enumerate(row, 1):
                    x = ws.cell(r, c)
                    if cell["f"]:
                        f = xl_prefix(cell["f"])
                        if cell.get("ov") and ";" in cell["ov"]:
                            n_down = len(cell["ov"].split(";")); n_across = max(len(p.split("~")) for p in cell["ov"].split(";"))
                            ref = f"{x.coordinate}:{col_letter(c - 1 + n_across - 1)}{r + n_down - 1}"
                            ws[x.coordinate] = ArrayFormula(ref, f)
                        else:
                            x.value = f
                    elif cell["v"] not in ("", None):
                        v = cell["v"]
                        if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v): v = datetime.date.fromisoformat(v)
                        x.value = v
                    if cell.get("fmt"): x.number_format = cell["fmt"].replace("₹", '"₹"')
                    for ch in cell.get("cls", ""):
                        if ch in fills: x.fill = PatternFill("solid", fgColor=fills[ch])
                        if ch in "hb": x.font = Font(bold=True, color="FFFFFF" if ch == "h" else "000000")
            for c in range(1, ws.max_column + 1):
                ws.column_dimensions[col_letter(c - 1)].width = 16
        wb.save(os.path.join(FILES, "excel", L.id + ".xlsx"))


if __name__ == "__main__":
    main(sys.argv[1:] or None)
